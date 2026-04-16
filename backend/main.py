"""
FastAPI server for WFM Schedule Generator.
"""
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
import shutil
import os
import json
import tempfile
from pathlib import Path

# Load .env file for local development
try:
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent.parent / ".env")
except ImportError:
    pass  # python-dotenv not installed, rely on system env vars

from data_parser import load_wfm_data, load_forecast_data
from solver import ScheduleSolver
from rule_engine import parse_regulations_with_llm, rules_to_json

app = FastAPI(
    title="WFM Schedule Generator",
    description="AI-powered workforce schedule generator using OR-Tools CP-SAT",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Store uploaded data and results in memory
app_state = {
    "wfm_data": None,
    "results": None,
    "uploaded_file": None,
    "forecast_uploaded": False,
    "forecast_file": None,
    "parsed_rules": None
}

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


@app.get("/api/health")
async def health_check():
    return {"status": "ok", "message": "WFM Schedule Generator API is running"}


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    """Upload a WFM Data Excel file."""
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "Please upload an Excel file (.xlsx)")
    
    filepath = UPLOAD_DIR / file.filename
    with open(filepath, "wb") as f:
        shutil.copyfileobj(file.file, f)
    
    try:
        wfm_data = load_wfm_data(str(filepath))
        app_state["wfm_data"] = wfm_data
        app_state["uploaded_file"] = str(filepath)
        
        # Parse regulations with GPT
        parsed_rules = parse_regulations_with_llm(wfm_data.regulations)
        app_state["parsed_rules"] = parsed_rules
        
        return {
            "message": "File uploaded and parsed successfully",
            "data": {
                "num_agents": len(wfm_data.agents),
                "num_dates": len(wfm_data.dates),
                "num_shift_codes": len(wfm_data.shift_codes),
                "num_leave_requests": len(wfm_data.leave_requests),
                "num_regulations": len(wfm_data.regulations),
                "agents": [
                    {
                        "id": a.id,
                        "name": a.name,
                        "role": a.role,
                        "channel": a.channel,
                        "gender": a.gender
                    }
                    for a in wfm_data.agents
                ],
                "dates": wfm_data.dates,
                "regulations": wfm_data.regulations,
                "parsed_rules": rules_to_json(parsed_rules),
                "leave_requests": [
                    {
                        "employee_id": lr.employee_id,
                        "employee_name": lr.employee_name,
                        "leave_type": lr.leave_type,
                        "date": lr.date.strftime('%Y-%m-%d')
                    }
                    for lr in wfm_data.leave_requests
                ]
            }
        }
    except Exception as e:
        raise HTTPException(500, f"Error parsing file: {str(e)}")


@app.post("/api/upload-forecast")
async def upload_forecast(file: UploadFile = File(...)):
    """Upload a Simulation Schedule / Forecast Excel file."""
    if app_state["wfm_data"] is None:
        raise HTTPException(400, "Please upload WFM data file first")
    
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "Please upload an Excel file (.xlsx)")
    
    filepath = UPLOAD_DIR / file.filename
    with open(filepath, "wb") as f:
        shutil.copyfileobj(file.file, f)
    
    try:
        wfm_data = app_state["wfm_data"]
        
        # Load forecast demand, mapping positionally to WFM dates
        # (forecast col 1 = WFM date 1, col 2 = WFM date 2, etc.)
        forecast = load_forecast_data(str(filepath), wfm_data.dates)
        wfm_data.forecast = forecast
        
        app_state["forecast_uploaded"] = True
        app_state["forecast_file"] = str(filepath)
        
        # Calculate summary stats
        wfm_dates = wfm_data.dates
        total_demand = 0
        for date_str in wfm_dates:
            reqs = forecast.requirements.get(date_str, {})
            total_demand += sum(reqs.values())
        avg_daily = total_demand // len(wfm_dates) if wfm_dates else 0
        
        return {
            "message": "Forecast data uploaded and parsed successfully",
            "data": {
                "num_dates": len(wfm_dates),
                "date_range": f"{wfm_dates[0]} to {wfm_dates[-1]}" if wfm_dates else "N/A",
                "dates": wfm_dates,
                "avg_daily_demand": avg_daily,
                "total_demand": total_demand,
                "leave_requests_in_range": len([
                    lr for lr in wfm_data.leave_requests
                    if lr.date.strftime('%Y-%m-%d') in wfm_dates
                ]),
                "leave_requests_total": len(wfm_data.leave_requests)
            }
        }
    except Exception as e:
        raise HTTPException(500, f"Error parsing forecast file: {str(e)}")


@app.post("/api/generate")
async def generate_schedule(time_limit: int = 60):
    """Generate an optimized schedule using CP-SAT solver."""
    if app_state["wfm_data"] is None:
        raise HTTPException(400, "No schedule data configured. Please set up agents and demand first.")
    
    try:
        solver = ScheduleSolver(app_state["wfm_data"])
        solver.build_model(parsed_rules=app_state.get("parsed_rules"))
        
        success = solver.solve(time_limit_seconds=time_limit)
        
        if not success:
            raise HTTPException(
                500, 
                "Solver could not find a feasible schedule. "
                "Try relaxing constraints or increasing the time limit."
            )
        
        results = solver.get_results_json()
        app_state["results"] = results
        
        return results
    
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error generating schedule: {str(e)}")


@app.post("/api/setup")
async def setup_schedule(data: dict):
    """Configure agents and shift demand from the form-based UI (no Excel required)."""
    from data_parser import build_wfm_from_form_data
    try:
        wfm_data = build_wfm_from_form_data(data)
        if not wfm_data.agents:
            raise HTTPException(400, "No agents provided")
        if not wfm_data.dates:
            raise HTTPException(400, "Invalid date range")

        app_state["wfm_data"] = wfm_data
        app_state["results"]   = None
        app_state["forecast_uploaded"] = True  # mark ready to generate

        return {
            "status": "ok",
            "agents": len(wfm_data.agents),
            "days":   len(wfm_data.dates),
            "dates":  wfm_data.dates,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Setup error: {str(e)}")


@app.get("/api/download-agent-template")
async def download_agent_template():
    """Generate and return an agent list Excel template for bulk import."""
    import openpyxl, io
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    wb = openpyxl.Workbook()

    # ── Styles ──────────────────────────────────────────────────────────────
    hdr_fill  = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    hdr_font  = Font(color="FFFFFF", bold=True, size=10)
    note_fill = PatternFill(start_color="e8eaf6", end_color="e8eaf6", fill_type="solid")
    note_font = Font(color="283593", italic=True, size=9)
    ref_fill  = PatternFill(start_color="e3f2fd", end_color="e3f2fd", fill_type="solid")
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'),  bottom=Side(style='thin')
    )
    center = Alignment(horizontal='center', vertical='center')

    def hdr(ws, row, col, value, width=None):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = center; c.border = thin
        return c

    def cell(ws, row, col, value, fill=None, bold=False):
        c = ws.cell(row=row, column=col, value=value)
        c.alignment = center; c.border = thin
        if fill: c.fill = fill
        if bold: c.font = Font(bold=True, size=10)
        return c

    # ── Sheet 1: Team List ───────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Team List"

    # Row 1: instructions
    ws.merge_cells('A1:H1')
    inst = ws['A1']
    inst.value = "Fill in your team details below. Do not change column headers. Refer to the 'Valid Values' sheet for accepted inputs."
    inst.font  = note_font
    inst.fill  = note_fill
    inst.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28

    # Row 2: column headers
    cols = ["No", "NIP", "Name", "Skill", "Channel", "Site", "Gender", "Religion"]
    for i, label in enumerate(cols, 1):
        hdr(ws, 2, i, label)

    # Sample data
    samples = [
        (1, "EMP001", "Ahmad Fauzi",    "Bahasa",  "Call",         "Jakarta",   "L", "Islam"),
        (2, "EMP002", "Siti Rahayu",    "Bahasa",  "Call",         "Jakarta",   "P", "Islam"),
        (3, "EMP003", "Budi Santoso",   "English", "Chat",         "Semarang",  "L", "Islam"),
        (4, "EMP004", "Dewi Lestari",   "Bahasa",  "Social Media", "Surabaya",  "P", "Kristen"),
        (5, "EMP005", "Rizky Pratama",  "English", "Email",        "Bandung",   "L", "Islam"),
    ]
    for r_offset, row_data in enumerate(samples):
        for c_offset, val in enumerate(row_data):
            cell(ws, 3 + r_offset, c_offset + 1, val)

    # Column widths
    for col_letter, width in zip("ABCDEFGH", [6, 12, 22, 10, 14, 14, 10, 12]):
        ws.column_dimensions[col_letter].width = width

    # ── Sheet 2: Valid Values ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Valid Values")
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 36

    ref_rows = [
        ("Column",      "Accepted Values"),
        ("Skill",       "Bahasa | English"),
        ("Channel",     "Social Media | Call | Email | Chat"),
        ("Site",        "Jakarta | Semarang | Surabaya | Bandung | Yogyakarta"),
        ("Gender",      "L  (Laki-laki / Male)   |   P  (Perempuan / Female)"),
        ("Religion",    "Islam | Kristen | Hindu | Buddha | Katolik"),
        ("NIP",         "Any unique employee ID string (e.g. EMP001)"),
        ("No",          "Row number — fill in or leave blank, will be ignored on import"),
    ]
    for r, (col_name, values) in enumerate(ref_rows, 1):
        c1 = ws2.cell(row=r, column=1, value=col_name)
        c2 = ws2.cell(row=r, column=2, value=values)
        if r == 1:
            c1.font = hdr_font; c1.fill = hdr_fill
            c2.font = hdr_font; c2.fill = hdr_fill
        else:
            c1.fill = ref_fill; c1.font = Font(bold=True, size=10)
            c2.fill = PatternFill(start_color="fafafa", end_color="fafafa", fill_type="solid")
        c1.border = thin; c2.border = thin
        c1.alignment = center
        c2.alignment = Alignment(horizontal='left', vertical='center')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Agent_List_Template.xlsx"}
    )


@app.post("/api/parse-agents")
async def parse_agents(file: UploadFile = File(...)):
    """Parse an agent list Excel file and return structured agent data."""
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "Please upload an Excel file (.xlsx or .xls)")

    try:
        import openpyxl, io

        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active

        # Find the header row — scan first 10 rows for one containing "name" or "nip"
        header_row_idx = None
        header_map = {}
        for row_idx, row in enumerate(ws.iter_rows(max_row=10, values_only=True), 1):
            row_lower = [str(v).strip().lower() if v is not None else '' for v in row]
            if 'name' in row_lower or 'nip' in row_lower:
                header_row_idx = row_idx
                for col_idx, val in enumerate(row_lower):
                    header_map[val] = col_idx
                break

        if header_row_idx is None:
            raise HTTPException(400, "Could not find a header row. Make sure your file has 'NIP' and 'Name' column headers.")

        # Normalise common aliases
        alias = {
            'employee id': 'nip', 'emp id': 'nip', 'id': 'nip',
            'full name': 'name', 'agent name': 'name',
            'skills': 'skill',
            'gender': 'gender', 'sex': 'gender',
            'religion': 'religion', 'agama': 'religion',
            'site': 'site', 'location': 'site',
            'channel': 'channel',
        }
        normalised = {}
        for raw_key, col_idx in header_map.items():
            key = alias.get(raw_key, raw_key)
            normalised[key] = col_idx

        valid_skills    = {'bahasa', 'english'}
        valid_channels  = {'social media', 'call', 'email', 'chat'}
        valid_sites     = {'jakarta', 'semarang', 'surabaya', 'bandung', 'yogyakarta'}
        valid_genders   = {'l', 'p', 'm', 'f'}
        valid_religions = {'islam', 'kristen', 'hindu', 'buddha', 'katolik'}

        def get_col(row, field):
            idx = normalised.get(field)
            if idx is None: return ''
            val = row[idx]
            return str(val).strip() if val is not None else ''

        agents = []
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            # Skip fully empty rows
            if all(v is None or str(v).strip() == '' for v in row):
                continue

            name = get_col(row, 'name')
            if not name:
                continue  # skip rows without a name

            nip      = get_col(row, 'nip')
            skill    = get_col(row, 'skill')
            channel  = get_col(row, 'channel')
            site     = get_col(row, 'site')
            gender   = get_col(row, 'gender').upper()
            religion = get_col(row, 'religion')

            # Normalise gender: M/male → L, F/female → P
            if gender in ('M', 'MALE', 'LAKI-LAKI', 'LAKI'):
                gender = 'L'
            elif gender in ('F', 'FEMALE', 'PEREMPUAN'):
                gender = 'P'

            # Validate / fallback to defaults
            if skill.lower() not in valid_skills:       skill    = 'Bahasa'
            if channel.lower() not in valid_channels:   channel  = 'Call'
            if site.lower() not in valid_sites:         site     = 'Jakarta'
            if gender not in ('L', 'P'):                gender   = 'L'
            if religion.lower() not in valid_religions: religion = 'Islam'

            agents.append({
                'nip':      nip,
                'name':     name,
                'skill':    skill,
                'channel':  channel,
                'site':     site,
                'gender':   gender,
                'religion': religion,
            })

        if not agents:
            raise HTTPException(400, "No valid agent rows found in the file. Check that the Name column has data.")

        return {"agents": agents, "count": len(agents)}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error parsing file: {str(e)}")


@app.get("/api/download-template")
async def download_template():
    """Generate and return an example WFM Data Excel file."""
    import openpyxl, io
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from datetime import datetime, timedelta

    wb = openpyxl.Workbook()

    # ── Styles ──────────────────────────────────────────────────────────────
    hdr_fill  = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    hdr_font  = Font(color="FFFFFF", bold=True, size=10)
    note_fill = PatternFill(start_color="e3f2fd", end_color="e3f2fd", fill_type="solid")
    note_font = Font(color="1565c0", italic=True, size=9)
    date_fill = PatternFill(start_color="fff8e1", end_color="fff8e1", fill_type="solid")
    thin      = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'),  bottom=Side(style='thin')
    )
    center    = Alignment(horizontal='center')

    def hdr(ws, row, col, value):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = hdr_fill; c.font = hdr_font; c.alignment = center; c.border = thin

    def cell(ws, row, col, value, fill=None):
        c = ws.cell(row=row, column=col, value=value)
        c.alignment = center; c.border = thin
        if fill: c.fill = fill
        return c

    # ── Sheet 1: Input ───────────────────────────────────────────────────────
    ws_input = wb.active
    ws_input.title = "Input"

    # Row 1 – headers
    for col, label in enumerate(["ID", "Name", "Role", "Channel", "Gender"], 1):
        hdr(ws_input, 1, col, label)

    # Generate 14 dates starting from the 1st of next month
    today     = datetime.today().replace(day=1)
    next_month= (today.replace(day=28) + timedelta(days=4)).replace(day=1)
    dates     = [next_month + timedelta(days=i) for i in range(14)]

    for i, dt in enumerate(dates):
        col = 6 + i
        hdr(ws_input, 1, col, dt.strftime('%d-%b'))
        c = ws_input.cell(row=2, column=col, value=dt)
        c.number_format = 'DD/MM/YYYY'
        c.fill = date_fill; c.alignment = center; c.border = thin

    # Row 2 col A–E: label hint
    ws_input.cell(row=2, column=1, value="← Date row →").font = Font(italic=True, color="888888", size=8)

    # Sample agents
    agents = [
        ("EMP001", "Ahmad Fauzi",        "Agent",    "Call", "L"),
        ("EMP002", "Siti Rahayu",         "Agent",    "Call", "P"),
        ("EMP003", "Budi Santoso",        "Senior",   "Call", "L"),
        ("EMP004", "Dewi Lestari",        "Agent",    "Call", "P"),
        ("EMP005", "Rizky Pratama",       "Agent",    "Call", "L"),
        ("EMP006", "Anisa Putri",         "Agent",    "Chat", "P"),
        ("EMP007", "Eko Wahyudi",         "TL",       "Call", "L"),
        ("EMP008", "Fitri Handayani",     "Agent",    "Call", "P"),
    ]
    for r, (emp_id, name, role, channel, gender) in enumerate(agents, 3):
        cell(ws_input, r, 1, emp_id)
        cell(ws_input, r, 2, name)
        cell(ws_input, r, 3, role)
        cell(ws_input, r, 4, channel)
        cell(ws_input, r, 5, gender)
        for i in range(14):
            cell(ws_input, r, 6 + i, "")

    # Column widths
    for col, w in zip("ABCDE", [10, 22, 10, 8, 8]):
        ws_input.column_dimensions[col].width = w
    for i in range(14):
        ws_input.column_dimensions[chr(70 + i)].width = 12

    # Note row
    note_row = len(agents) + 4
    note = ws_input.cell(row=note_row, column=1,
        value="Gender codes: L = Laki-laki (Male), P = Perempuan (Female)  |  Leave date cells blank — the solver assigns shifts automatically")
    note.font = note_font; note.fill = note_fill
    ws_input.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)

    # ── Sheet 2: Shift Code ──────────────────────────────────────────────────
    ws_shift = wb.create_sheet("Shift Code")
    for col, label in enumerate(["Code", "Time Range", "Description"], 1):
        hdr(ws_shift, 1, col, label)

    shift_codes = [
        ("P1",  "06:00-15:00", "Morning early"),
        ("P2",  "07:00-16:00", "Morning"),
        ("P3",  "08:00-17:00", "Morning standard"),
        ("P4",  "09:00-18:00", "Morning late"),
        ("P10", "10:00-19:00", "Mid-morning"),
        ("S1",  "11:00-20:00", "Afternoon"),
        ("S2",  "12:00-21:00", "Afternoon"),
        ("S4",  "13:00-22:00", "Afternoon late"),
        ("S7",  "15:00-00:00", "Evening"),
        ("M3",  "21:00-06:00", "Night"),
        ("M1",  "22:00-07:00", "Night late"),
        ("OFF", "",            "Rest day"),
    ]
    for r, (code, time_range, desc) in enumerate(shift_codes, 2):
        cell(ws_shift, r, 1, code)
        cell(ws_shift, r, 2, time_range)
        cell(ws_shift, r, 3, desc)

    for col, w in zip("ABC", [8, 14, 20]):
        ws_shift.column_dimensions[col].width = w

    # ── Sheet 3: Regulation ──────────────────────────────────────────────────
    ws_reg = wb.create_sheet("Regulation")
    hdr(ws_reg, 1, 2, "Regulation Text")

    regulations = [
        "Maximum 6 consecutive working days",
        "Maximum 2 consecutive days off",
        "No shift jumping — agents must stay on the same shift type unless they take an OFF day",
        "Female agents (gender P) may not be assigned to night shifts (21:00 or later)",
    ]
    for r, reg in enumerate(regulations, 2):
        c = ws_reg.cell(row=r, column=2, value=reg)
        c.border = thin

    ws_reg.column_dimensions['B'].width = 70

    # ── Sheet 4: Leave_Req OFF Tracker ───────────────────────────────────────
    ws_leave = wb.create_sheet("Leave_Req OFF Tracker")
    for col, label in enumerate(["Employee ID", "Name", "Note", "Note2", "Leave Type", "Date"], 1):
        hdr(ws_leave, 1, col, label)

    leave_data = [
        ("EMP001", "Ahmad Fauzi",    "", "", "Off",   dates[3]),
        ("EMP002", "Siti Rahayu",    "", "", "Leave", dates[5]),
        ("EMP003", "Budi Santoso",   "", "", "Off",   dates[6]),
        ("EMP004", "Dewi Lestari",   "", "", "Leave", dates[7]),
        ("EMP005", "Rizky Pratama",  "", "", "Off",   dates[8]),
    ]
    for r, (emp_id, name, n1, n2, ltype, dt) in enumerate(leave_data, 2):
        cell(ws_leave, r, 1, emp_id)
        cell(ws_leave, r, 2, name)
        cell(ws_leave, r, 3, n1)
        cell(ws_leave, r, 4, n2)
        cell(ws_leave, r, 5, ltype)
        c = ws_leave.cell(row=r, column=6, value=dt)
        c.number_format = 'DD/MM/YYYY'
        c.alignment = center; c.border = thin

    for col, w in zip("ABCDEF", [12, 22, 8, 8, 12, 14]):
        ws_leave.column_dimensions[col].width = w

    # ── Save & return ────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=WFM_Data_Template.xlsx"}
    )


@app.get("/api/results")
async def get_results():
    """Get the most recent schedule results."""
    if app_state["results"] is None:
        raise HTTPException(404, "No schedule has been generated yet")
    
    return app_state["results"]


@app.get("/api/export")
async def export_excel():
    """Export the generated schedule as an Excel file."""
    if app_state["results"] is None:
        raise HTTPException(404, "No schedule has been generated yet")
    
    try:
        import openpyxl, io
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Generated Schedule"
        
        results = app_state["results"]
        dates = results["dates"]
        agents = results["agents"]
        
        # Styles
        header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        off_fill = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")
        morning_fill = PatternFill(start_color="e3f2fd", end_color="e3f2fd", fill_type="solid")
        evening_fill = PatternFill(start_color="fff3e0", end_color="fff3e0", fill_type="solid")
        leave_fill = PatternFill(start_color="fce4ec", end_color="fce4ec", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Headers
        headers = ["ID", "Name", "Role", "Channel", "Gender"] + dates
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # Data rows
        for row_idx, agent in enumerate(agents, 2):
            ws.cell(row=row_idx, column=1, value=agent['id']).border = thin_border
            ws.cell(row=row_idx, column=2, value=agent['name']).border = thin_border
            ws.cell(row=row_idx, column=3, value=agent['role']).border = thin_border
            ws.cell(row=row_idx, column=4, value=agent['channel']).border = thin_border
            ws.cell(row=row_idx, column=5, value=agent['gender']).border = thin_border
            
            for col_idx, date in enumerate(dates, 6):
                shift = agent['schedule'].get(date, 'OFF')
                cell = ws.cell(row=row_idx, column=col_idx, value=shift)
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border
                
                if shift == 'OFF':
                    cell.fill = off_fill
                elif shift == '03:00':
                    cell.fill = morning_fill
                elif shift == '20:00':
                    cell.fill = evening_fill
                elif shift in ('Leave', 'Resign'):
                    cell.fill = leave_fill
        
        # Auto-width
        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 15)
        
        # Coverage stats sheet
        ws2 = wb.create_sheet("Coverage Stats")
        coverage = results.get("coverage", {})
        
        cov_headers = ["Date", "03:00 Shift", "20:00 Shift", "Total Working", "OFF", "Leave", "Demand (avg)"]
        for col, header in enumerate(cov_headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
        
        for row_idx, date in enumerate(dates, 2):
            stats = coverage.get(date, {})
            ws2.cell(row=row_idx, column=1, value=date).border = thin_border
            ws2.cell(row=row_idx, column=2, value=stats.get('shift_03_count', 0)).border = thin_border
            ws2.cell(row=row_idx, column=3, value=stats.get('shift_20_count', 0)).border = thin_border
            ws2.cell(row=row_idx, column=4, value=stats.get('total_working', 0)).border = thin_border
            ws2.cell(row=row_idx, column=5, value=stats.get('off_count', 0)).border = thin_border
            ws2.cell(row=row_idx, column=6, value=stats.get('leave_count', 0)).border = thin_border
            ws2.cell(row=row_idx, column=7, value=stats.get('forecast_demand', 0)).border = thin_border
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=WFM_Generated_Schedule.xlsx"}
        )
    except Exception as e:
        raise HTTPException(500, f"Error exporting: {str(e)}")


# ─── Serve frontend static files ───
frontend_dir = Path(__file__).parent.parent / "frontend"
if frontend_dir.exists():
    app.mount("/", StaticFiles(directory=str(frontend_dir), html=True), name="frontend")


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
